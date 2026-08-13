import io
from datetime import date, datetime
from decimal import Decimal

from django.contrib.contenttypes.models import ContentType
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader

from apps.accounts.models import Role, User, UserLocationAssignment, UserRoleAssignment
from apps.analytics.exports import _excel_value, safe_spreadsheet_value
from apps.analytics.models import DataExportAudit
from apps.documents.models import EvidenceDocument, EvidenceLink
from apps.locations.models import Province, Tikina, Village
from apps.population.models import Household
from apps.projects.models import IVDPProject
from apps.reporting.models import ReportingPeriod, TNKReport
from apps.reporting.services import create_report
from apps.workflow.models import ApprovalAction


class ExportQualityTests(TestCase):
    def setUp(self):
        self.province = Province.objects.create(code="PDF", name_en="Naitāsiri Province")
        self.tikina = Tikina.objects.create(province=self.province, code="PDF-T", name_en="Waimaro Tikina")
        self.period = ReportingPeriod.objects.create(
            year=2026,
            quarter=2,
            start_date=date(2026, 4, 1),
            end_date=date(2026, 6, 30),
            submission_due_date=date(2026, 7, 15),
            is_open=True,
        )
        self.author = self.make_user("pdf-author", Role.Codes.TURAGA_NI_KORO)
        self.analyst = self.make_user("pdf-analyst", Role.Codes.READ_ONLY_ANALYST)
        UserLocationAssignment.objects.create(user=self.author, province=self.province)
        UserLocationAssignment.objects.create(user=self.analyst, province=self.province)
        self.reports = []
        for number, name in enumerate(
            (
                "Vunivau ā ē ī ō ū",
                "Nakorosule with a deliberately long village name used to confirm wrapped table text",
                "=HYPERLINK(\"https://invalid.example\")",
                "Nauluvatu",
            ),
            1,
        ):
            village = Village.objects.create(tikina=self.tikina, code=f"PDF-{number}", name_en=name)
            report = create_report(village=village, reporting_period=self.period, prepared_by=self.author)
            TNKReport.objects.filter(pk=report.pk).update(
                status=TNKReport.Status.APPROVED,
                completeness_percentage=Decimal("100"),
                data_quality_score=Decimal("87.50"),
                approved_at=timezone.now(),
            )
            report.refresh_from_db()
            ApprovalAction.objects.create(
                report=report,
                user=self.author,
                user_full_name="Jone Vakaloloma",
                user_role="Roko Tui",
                action_type="approve",
                from_status=TNKReport.Status.UNDER_PROVINCIAL_REVIEW,
                to_status=TNKReport.Status.APPROVED,
                comment=(
                    "Vakadeitaki ni ripote. "
                    "This deliberately long approval comment verifies that table content wraps across lines without clipping. "
                    * 3
                ),
                digital_acknowledgement=True,
            )
            self.reports.append(report)
        self.project = IVDPProject.objects.create(
            project_code="IVDP-UNICODE",
            village=self.reports[0].village,
            project_name="Vale ni Bose kei na wai savasava - " + "a very long project title " * 5,
            project_category="water",
            problem_being_addressed="Reliable village water access",
            priority="high",
            planned_start_date=date(2026, 4, 2),
            planned_end_date=date(2026, 6, 20),
            approved_budget=Decimal("123456.78"),
            currency_code="FJD",
            project_status="active",
            physical_progress_percentage=Decimal("55.50"),
            financial_progress_percentage=Decimal("42.25"),
        )
        Household.objects.create(
            household_code="SECRET-1",
            village=self.reports[0].village,
            household_head_name="RESTRICTED PERSON MUST NOT EXPORT",
            household_size=4,
            effective_from=date(2026, 1, 1),
        )

    def make_user(self, username, role_code):
        user = User.objects.create_user(username=username, password="test-password")
        role, _ = Role.objects.get_or_create(code=role_code, defaults={"name": Role(code=role_code).get_code_display()})
        UserRoleAssignment.objects.create(user=user, role=role)
        return user

    @override_settings(MEDIA_ROOT="tmp/test-phase-f-media")
    def test_pdf_supports_unicode_pagination_wrapping_history_projects_and_evidence(self):
        document = EvidenceDocument.objects.create(
            title="iVakadinadina ni Bose ā",
            document_type="meeting_minutes",
            file=SimpleUploadedFile("minutes.pdf", b"%PDF-1.4 evidence", content_type="application/pdf"),
            original_filename="minutes.pdf",
            file_size=17,
            mime_type="application/pdf",
            checksum="0" * 64,
            confidentiality_level="restricted",
            uploaded_by=self.author,
        )
        EvidenceLink.objects.create(
            document=document,
            content_type=ContentType.objects.get_for_model(TNKReport),
            object_id=self.reports[0].pk,
        )
        self.client.force_login(self.author)
        response = self.client.get(reverse("analytics:export", args=("pdf",)), {"year": "2026", "quarter": "2"})
        self.assertEqual(response.status_code, 200)
        reader = PdfReader(io.BytesIO(response.content))
        self.assertGreaterEqual(len(reader.pages), 5)
        self.assertEqual(reader.metadata.title, "TNK Insight - Quarterly Report Summary")
        page_text = [page.extract_text() for page in reader.pages]
        text = "\n".join(page_text)
        self.assertIn("Vunivau ā ē ī ō ū", text)
        self.assertIn("Approval history", text)
        self.assertIn("IVDP project summary", text)
        self.assertIn("Vale ni Bose kei na wai savasava", text)
        self.assertIn("Evidence references", text)
        self.assertIn("iVakadinadina ni Bose ā", text)
        self.assertIn(f"Page 1 of {len(reader.pages)}", page_text[0])
        self.assertIn(f"Page {len(reader.pages)} of {len(reader.pages)}", page_text[-1])
        self.assertNotIn("RESTRICTED PERSON MUST NOT EXPORT", text)

    def test_analyst_pdf_contains_scoped_summary_but_no_named_detail(self):
        self.client.force_login(self.analyst)
        response = self.client.get(reverse("analytics:export", args=("pdf",)))
        reader = PdfReader(io.BytesIO(response.content))
        text = "\n".join(page.extract_text() for page in reader.pages)
        self.assertIn("Naitāsiri Province", text)
        self.assertNotIn("Approval history", text)
        self.assertNotIn("Jone Vakaloloma", text)
        self.assertNotIn("Vale ni Bose kei na wai savasava", text)

    def test_excel_uses_real_types_unicode_readable_headings_and_clear_filters(self):
        self.client.force_login(self.author)
        response = self.client.get(
            reverse("analytics:export", args=("xlsx",)),
            {"year": "2026", "quarter": "2", "reason": "Quarterly review"},
        )
        workbook = load_workbook(io.BytesIO(response.content), data_only=False)
        reports = workbook["Reports"]
        self.assertEqual(reports.freeze_panes, "A2")
        self.assertTrue(reports.auto_filter.ref)
        self.assertEqual(reports["A1"].value, "Village")
        values = [reports.cell(row=row, column=1).value for row in range(2, reports.max_row + 1)]
        self.assertIn("Vunivau ā ē ī ō ū", values)
        self.assertIn("'=HYPERLINK(\"https://invalid.example\")", values)
        self.assertIsInstance(reports["D2"].value, int)
        self.assertIsInstance(reports["F2"].value, (date, datetime))
        self.assertIsInstance(reports["I2"].value, (int, float))
        information = workbook["Export information"]
        self.assertEqual(information["B4"].value, "Quarterly review")
        self.assertEqual(information["B9"].value, "2")
        projects = workbook["Projects"]
        self.assertEqual(projects["D2"].value, self.project.project_name)
        self.assertIsInstance(projects["F2"].value, (date, datetime))
        self.assertIsInstance(projects["H2"].value, (int, float))
        self.assertEqual(projects["H2"].value, 123456.78)
        self.assertNotIn("RESTRICTED PERSON MUST NOT EXPORT", str(workbook.sheetnames) + str(values))

    def test_csv_is_utf8_bom_unicode_typed_text_and_formula_safe(self):
        self.client.force_login(self.author)
        response = self.client.get(reverse("analytics:export", args=("csv",)))
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))
        text = response.content.decode("utf-8-sig")
        self.assertIn("Vunivau ā ē ī ō ū", text)
        self.assertIn("'=HYPERLINK", text)
        self.assertNotIn("RESTRICTED PERSON MUST NOT EXPORT", text)
        self.assertEqual(safe_spreadsheet_value("+SUM(A1:A2)"), "'+SUM(A1:A2)")
        self.assertIsInstance(_excel_value(Decimal("123.45")), float)

    def test_each_export_is_audited_with_applied_filters_and_count(self):
        self.client.force_login(self.author)
        for file_format in ("csv", "xlsx", "pdf"):
            response = self.client.get(
                reverse("analytics:export", args=(file_format,)),
                {"year": "2026", "quarter": "2", "reason": "Governance review"},
            )
            self.assertEqual(response.status_code, 200)
        audits = DataExportAudit.objects.order_by("pk")
        self.assertEqual(audits.count(), 3)
        for audit in audits:
            self.assertEqual(audit.number_of_records, 4)
            self.assertEqual(audit.filters_used, {"year": "2026", "quarter": "2"})
            self.assertEqual(audit.export_reason, "Governance review")
