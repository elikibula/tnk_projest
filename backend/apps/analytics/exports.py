import csv
import io
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.contrib.contenttypes.models import ContentType
from django.core.exceptions import ImproperlyConfigured, PermissionDenied
from django.http import HttpResponse
from django.utils import timezone
from django.utils.html import strip_tags
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    KeepTogether,
    LongTable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from apps.accounts.selectors import villages_for_user
from apps.audit.services import record_event
from apps.core.security import can_export_report_summary, can_view_document, can_view_entry, can_view_report
from apps.documents.models import EvidenceDocument
from apps.reporting.models import ReportAmendment
from apps.reporting.section_entries import entry_queryset
from apps.reporting.section_registry import get_entry_config

from .models import DataExportAudit


HEADERS = (
    "Village",
    "Tikina",
    "Province",
    "Year",
    "Quarter",
    "Period start",
    "Period end",
    "Status",
    "Completion (%)",
    "Data quality score",
    "Approved date",
)
FORMULA_PREFIXES = ("=", "+", "-", "@")
PDF_TITLE = "TNK Insight - Quarterly Report Summary"
NAVY = colors.HexColor("#173A4D")
TEAL = colors.HexColor("#0F766E")
PALE_TEAL = colors.HexColor("#E6F4F1")
PALE_GREY = colors.HexColor("#F3F6F8")
TEXT_GREY = colors.HexColor("#475569")


@dataclass(frozen=True)
class ExportDataset:
    reports: tuple
    rows: tuple[tuple, ...]
    filters: dict
    reason: str
    generated_at: object


def safe_spreadsheet_value(value):
    """Prevent user-entered text from being evaluated as a spreadsheet formula."""
    if isinstance(value, str) and value.lstrip().startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def _safe_text(value):
    return strip_tags(str(value or "")).replace("\x00", "")


def rows(reports):
    return [
        (
            safe_spreadsheet_value(report.village.name_en),
            safe_spreadsheet_value(report.village.tikina.name_en),
            safe_spreadsheet_value(report.village.tikina.province.name_en),
            report.reporting_period.year,
            report.reporting_period.quarter,
            report.reporting_period.start_date,
            report.reporting_period.end_date,
            report.get_status_display(),
            report.completeness_percentage,
            report.data_quality_score,
            report.approved_at.date() if report.approved_at else None,
        )
        for report in reports
    ]


def authorised_reports(user, reports):
    if not can_export_report_summary(user):
        raise PermissionDenied("Your role cannot export TNK analytics.")
    if hasattr(reports, "filter"):
        reports = reports.filter(village__in=villages_for_user(user))
        return reports.select_related("village__tikina__province", "reporting_period", "prepared_by").prefetch_related(
            "approval_actions", "amendments__requested_by", "amendments__approved_by"
        )
    village_ids = set(villages_for_user(user).values_list("pk", flat=True))
    return [report for report in reports if report.village_id in village_ids]


def build_export_dataset(user, reports, filters, reason):
    scoped = tuple(authorised_reports(user, reports))
    return ExportDataset(
        reports=scoped,
        rows=tuple(rows(scoped)),
        filters=dict(filters),
        reason=_safe_text(reason),
        generated_at=timezone.now(),
    )


def audit(user, fmt, count, filters, reason):
    export = DataExportAudit.objects.create(
        generated_by=user,
        report_type="tnk_report_summary",
        filters_used=filters,
        number_of_records=count,
        export_reason=reason,
        format=fmt,
    )
    record_event(actor=user, action="export.created", instance=export, summary=f"Created {fmt} export")


def _response(content, content_type, filename):
    response = HttpResponse(content, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["X-Content-Type-Options"] = "nosniff"
    return response


def csv_export(user, reports, filters, reason):
    dataset = build_export_dataset(user, reports, filters, reason)
    output = io.StringIO(newline="")
    output.write("\ufeff")
    writer = csv.writer(output)
    writer.writerow(HEADERS)
    writer.writerows(dataset.rows)
    audit(user, "csv", len(dataset.rows), dataset.filters, dataset.reason)
    return _response(output.getvalue(), "text/csv", "tnk-reports.csv")


def _excel_value(value):
    value = safe_spreadsheet_value(value)
    if isinstance(value, Decimal):
        return float(value)
    return value


def _format_filter_value(value):
    if value in (None, ""):
        return "All"
    return safe_spreadsheet_value(_safe_text(value))


def xlsx_export(user, reports, filters, reason):
    dataset = build_export_dataset(user, reports, filters, reason)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reports"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(1, len(dataset.rows) + 1)}"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in dataset.rows:
        sheet.append([_excel_value(value) for value in row])
    for row_number in range(2, sheet.max_row + 1):
        for column in (6, 7, 11):
            sheet.cell(row=row_number, column=column).number_format = "dd mmm yyyy"
        for column in (9, 10):
            sheet.cell(row=row_number, column=column).number_format = "0.00"
    widths = (30, 24, 24, 10, 10, 15, 15, 24, 18, 20, 16)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 32

    information = workbook.create_sheet("Export information")
    information.append(("TNK Insight export information", ""))
    information.append(("Generated", dataset.generated_at.replace(tzinfo=None)))
    information.append(("Records", len(dataset.rows)))
    information.append(("Reason", safe_spreadsheet_value(dataset.reason)))
    for key in ("province", "tikina", "village", "year", "quarter", "status"):
        information.append((key.replace("_", " ").title(), _format_filter_value(dataset.filters.get(key))))
    information["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    information["A1"].fill = PatternFill("solid", fgColor="173A4D")
    information["B1"].fill = PatternFill("solid", fgColor="173A4D")
    information.column_dimensions["A"].width = 24
    information.column_dimensions["B"].width = 55
    information["B2"].number_format = "dd mmm yyyy hh:mm"
    information.freeze_panes = "A2"
    project_rows = []
    for report in dataset.reports:
        for project in _projects_for_report(user, report):
            project_rows.append(
                (
                    report.village.name_en,
                    str(report.reporting_period),
                    project.project_code,
                    project.project_name,
                    project.get_project_status_display() if hasattr(project, "get_project_status_display") else project.project_status,
                    project.planned_start_date,
                    project.planned_end_date,
                    project.approved_budget,
                    project.currency_code,
                    project.physical_progress_percentage,
                    project.financial_progress_percentage,
                )
            )
    if project_rows:
        projects = workbook.create_sheet("Projects")
        project_headers = (
            "Village",
            "Reporting period",
            "Project code",
            "Project name",
            "Status",
            "Planned start",
            "Planned end",
            "Approved budget",
            "Currency",
            "Physical progress (%)",
            "Financial progress (%)",
        )
        projects.append(project_headers)
        for row in project_rows:
            projects.append([_excel_value(value) for value in row])
        projects.freeze_panes = "A2"
        projects.auto_filter.ref = f"A1:K{projects.max_row}"
        for cell in projects[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173A4D")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row_number in range(2, projects.max_row + 1):
            for column in (6, 7):
                projects.cell(row=row_number, column=column).number_format = "dd mmm yyyy"
            projects.cell(row=row_number, column=8).number_format = '#,##0.00 "FJD"'
            for column in (10, 11):
                projects.cell(row=row_number, column=column).number_format = "0.00"
        project_widths = (28, 18, 18, 55, 20, 15, 15, 20, 12, 22, 22)
        for column, width in enumerate(project_widths, 1):
            projects.column_dimensions[get_column_letter(column)].width = width
    workbook.properties.title = PDF_TITLE
    workbook.properties.subject = "Scoped TNK quarterly report export"

    output = io.BytesIO()
    workbook.save(output)
    audit(user, "xlsx", len(dataset.rows), dataset.filters, dataset.reason)
    return _response(
        output.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "tnk-reports.xlsx",
    )


def _font_candidates(bold=False):
    configured = getattr(settings, "TNK_PDF_FONT_BOLD_PATH" if bold else "TNK_PDF_FONT_PATH", "")
    names = (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    )
    reportlab_font = Path(__file__).resolve().parents[3] / "unused"
    try:
        import reportlab

        reportlab_font = Path(reportlab.__file__).resolve().parent / "fonts" / ("VeraBd.ttf" if bold else "Vera.ttf")
    except ImportError:
        pass
    return tuple(path for path in (configured, *names, str(reportlab_font)) if path)


def _register_pdf_fonts():
    regular = next((path for path in _font_candidates() if Path(path).is_file()), None)
    bold = next((path for path in _font_candidates(bold=True) if Path(path).is_file()), None)
    if regular is None or bold is None:
        raise ImproperlyConfigured(
            "No Unicode PDF font is available. Install DejaVu Sans or configure TNK_PDF_FONT_PATH and TNK_PDF_FONT_BOLD_PATH."
        )
    if "TNKUnicode" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("TNKUnicode", regular))
    if "TNKUnicode-Bold" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("TNKUnicode-Bold", bold))
    pdfmetrics.registerFontFamily(
        "TNKUnicode",
        normal="TNKUnicode",
        bold="TNKUnicode-Bold",
        italic="TNKUnicode",
        boldItalic="TNKUnicode-Bold",
    )
    return "TNKUnicode", "TNKUnicode-Bold"


class PageNumberCanvasMixin:
    """Mixin factory state used by the reportlab canvas class in `_pdf_canvasmaker`."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        page_count = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self._draw_page_frame(page_count)
            super().showPage()
        super().save()

    def _draw_page_frame(self, page_count):
        self.saveState()
        self.setTitle(PDF_TITLE)
        self.setAuthor("TNK Insight")
        self.setFont("TNKUnicode", 8)
        self.setFillColor(TEXT_GREY)
        self.drawString(18 * mm, 12 * mm, "TNK Insight - Official scoped export")
        self.drawRightString(A4[0] - 18 * mm, 12 * mm, f"Page {self._pageNumber} of {page_count}")
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.line(18 * mm, 16 * mm, A4[0] - 18 * mm, 16 * mm)
        self.restoreState()


def _pdf_canvasmaker():
    from reportlab.pdfgen.canvas import Canvas

    return type("TNKNumberedCanvas", (PageNumberCanvasMixin, Canvas), {})


def _pdf_styles(regular, bold):
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "TNKTitle", parent=base["Title"], fontName=bold, fontSize=20, leading=25, textColor=NAVY, alignment=TA_LEFT
        ),
        "subtitle": ParagraphStyle(
            "TNKSubtitle", parent=base["Normal"], fontName=regular, fontSize=9, leading=13, textColor=TEXT_GREY
        ),
        "heading": ParagraphStyle(
            "TNKHeading", parent=base["Heading2"], fontName=bold, fontSize=13, leading=17, textColor=NAVY, spaceAfter=5
        ),
        "subheading": ParagraphStyle(
            "TNKSubheading", parent=base["Heading3"], fontName=bold, fontSize=10, leading=13, textColor=TEAL, spaceAfter=4
        ),
        "body": ParagraphStyle("TNKBody", parent=base["BodyText"], fontName=regular, fontSize=8.5, leading=12),
        "small": ParagraphStyle("TNKSmall", parent=base["BodyText"], fontName=regular, fontSize=7.5, leading=10),
        "cell": ParagraphStyle("TNKCell", parent=base["BodyText"], fontName=regular, fontSize=7, leading=9),
        "cell_bold": ParagraphStyle("TNKCellBold", parent=base["BodyText"], fontName=bold, fontSize=7, leading=9),
        "center": ParagraphStyle(
            "TNKCenter", parent=base["BodyText"], fontName=regular, fontSize=7, leading=9, alignment=TA_CENTER
        ),
    }


def _paragraph(value, style):
    from xml.sax.saxutils import escape

    return Paragraph(escape(_safe_text(value)) or "-", style)


def _table(data, widths, *, header=True, repeat_rows=1):
    table = LongTable(data, colWidths=widths, repeatRows=repeat_rows if header else 0, hAlign="LEFT")
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1 if header else 0), (-1, -1), (colors.white, PALE_GREY)),
    ]
    if header:
        commands.extend(
            [
                ("BACKGROUND", (0, 0), (-1, 0), TEAL),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ]
        )
    table.setStyle(TableStyle(commands))
    return table


def _filter_summary(dataset, styles):
    active = [f"{key.replace('_', ' ').title()}: {_safe_text(value)}" for key, value in dataset.filters.items() if value]
    text = " | ".join(active) if active else "All reports within the user's assigned scope"
    reason = dataset.reason or "Not supplied"
    return Table(
        [
            [_paragraph("Export scope", styles["cell_bold"]), _paragraph(text, styles["cell"])],
            [_paragraph("Generated", styles["cell_bold"]), _paragraph(timezone.localtime(dataset.generated_at).strftime("%d %b %Y %H:%M %Z"), styles["cell"])],
            [_paragraph("Reason", styles["cell_bold"]), _paragraph(reason, styles["cell"])],
        ],
        colWidths=(35 * mm, 140 * mm),
        style=TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), PALE_TEAL),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        ),
    )


def _summary_table(dataset, styles):
    header = [
        _paragraph(label, styles["cell_bold"])
        for label in ("Village", "Tikina", "Province", "Period", "Status", "Completion", "Quality")
    ]
    data = [header]
    for report in dataset.reports:
        data.append(
            [
                _paragraph(report.village.name_en, styles["cell"]),
                _paragraph(report.village.tikina.name_en, styles["cell"]),
                _paragraph(report.village.tikina.province.name_en, styles["cell"]),
                _paragraph(str(report.reporting_period), styles["center"]),
                _paragraph(report.get_status_display(), styles["cell"]),
                _paragraph(f"{report.completeness_percentage}%", styles["center"]),
                _paragraph(report.data_quality_score if report.data_quality_score is not None else "Not scored", styles["center"]),
            ]
        )
    if len(data) == 1:
        data.append([_paragraph("No reports matched the selected filters.", styles["cell"]), "", "", "", "", "", ""])
    return _table(data, (38 * mm, 28 * mm, 28 * mm, 21 * mm, 25 * mm, 18 * mm, 18 * mm))


def _report_metadata(report, styles):
    items = (
        ("Village", report.village.name_en),
        ("Tikina", report.village.tikina.name_en),
        ("Province", report.village.tikina.province.name_en),
        ("Reporting period", str(report.reporting_period)),
        ("Period dates", f"{report.reporting_period.start_date:%d %b %Y} to {report.reporting_period.end_date:%d %b %Y}"),
        ("Report status", report.get_status_display()),
        ("Completion", f"{report.completeness_percentage}%"),
        ("Data-quality score", report.data_quality_score if report.data_quality_score is not None else "Not scored"),
    )
    data = []
    for index in range(0, len(items), 2):
        left, right = items[index], items[index + 1]
        data.append(
            [
                _paragraph(left[0], styles["cell_bold"]),
                _paragraph(left[1], styles["cell"]),
                _paragraph(right[0], styles["cell_bold"]),
                _paragraph(right[1], styles["cell"]),
            ]
        )
    table = Table(data, colWidths=(30 * mm, 57 * mm, 30 * mm, 58 * mm))
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), PALE_TEAL),
                ("BACKGROUND", (2, 0), (2, -1), PALE_TEAL),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def _approval_table(report, styles):
    actions = list(report.approval_actions.all().order_by("acted_at", "pk"))
    data = [[_paragraph(label, styles["cell_bold"]) for label in ("Action", "From / to", "Officer", "Role", "Date", "Comment")]]
    for action in actions:
        data.append(
            [
                _paragraph(action.get_action_type_display(), styles["cell"]),
                _paragraph(f"{action.from_status.replace('_', ' ')} -> {action.to_status.replace('_', ' ')}", styles["cell"]),
                _paragraph(action.user_full_name or action.user.username, styles["cell"]),
                _paragraph(action.user_role, styles["cell"]),
                _paragraph(timezone.localtime(action.acted_at).strftime("%d %b %Y %H:%M"), styles["cell"]),
                _paragraph(action.comment or "-", styles["cell"]),
            ]
        )
    if len(data) == 1:
        data.append([_paragraph("No approval actions recorded.", styles["cell"]), "", "", "", "", ""])
    return _table(data, (21 * mm, 35 * mm, 27 * mm, 25 * mm, 28 * mm, 39 * mm))


def _projects_for_report(user, report):
    config = get_entry_config("ivdp_projects", "project")
    if config is None or not can_view_entry(user, "ivdp_projects", config.model, report.village):
        return ()
    return entry_queryset(config, report).order_by("project_code", "pk")


def _project_table(user, report, styles):
    projects = list(_projects_for_report(user, report))
    if not projects:
        return None
    data = [
        [
            _paragraph(label, styles["cell_bold"])
            for label in ("Code", "Project", "Status", "Planned dates", "Budget", "Physical / financial")
        ]
    ]
    for project in projects:
        planned_dates = "-"
        if project.planned_start_date or project.planned_end_date:
            start = project.planned_start_date.strftime("%d %b %Y") if project.planned_start_date else "Not set"
            end = project.planned_end_date.strftime("%d %b %Y") if project.planned_end_date else "Not set"
            planned_dates = f"{start} to {end}"
        budget = f"{project.currency_code} {project.approved_budget:,.2f}" if project.approved_budget is not None else "Not set"
        data.append(
            [
                _paragraph(project.project_code, styles["cell"]),
                _paragraph(project.project_name, styles["cell"]),
                _paragraph(
                    project.get_project_status_display() if hasattr(project, "get_project_status_display") else project.project_status,
                    styles["cell"],
                ),
                _paragraph(planned_dates, styles["cell"]),
                _paragraph(budget, styles["cell"]),
                _paragraph(
                    f"{project.physical_progress_percentage}% / {project.financial_progress_percentage}%",
                    styles["cell"],
                ),
            ]
        )
    return _table(data, (18 * mm, 52 * mm, 25 * mm, 32 * mm, 25 * mm, 23 * mm))


def _amendment_table(report, styles):
    amendments = report.amendments.filter(status=ReportAmendment.Status.APPROVED).order_by("amendment_number")
    data = [[_paragraph(label, styles["cell_bold"]) for label in ("No.", "Reason", "Requested by", "Approved by", "Approval date")]]
    for amendment in amendments:
        data.append(
            [
                _paragraph(amendment.amendment_number, styles["center"]),
                _paragraph(amendment.reason, styles["cell"]),
                _paragraph(amendment.requested_by.get_full_name() or amendment.requested_by.username, styles["cell"]),
                _paragraph(amendment.approved_by.get_full_name() or amendment.approved_by.username, styles["cell"]),
                _paragraph(timezone.localtime(amendment.approved_at).strftime("%d %b %Y") if amendment.approved_at else "-", styles["cell"]),
            ]
        )
    return _table(data, (12 * mm, 70 * mm, 34 * mm, 34 * mm, 25 * mm)) if len(data) > 1 else None


def _evidence_for_reports(user, reports):
    report_ids = [report.pk for report in reports]
    if not report_ids:
        return {}
    content_type = ContentType.objects.get_for_model(reports[0])
    documents = EvidenceDocument.objects.filter(
        links__content_type=content_type,
        links__object_id__in=report_ids,
    ).prefetch_related("links__content_type").distinct()
    allowed = {document.pk: document for document in documents if can_view_document(user, document)}
    result = {report_id: [] for report_id in report_ids}
    for document in allowed.values():
        for link in document.links.all():
            if link.content_type_id == content_type.pk and link.object_id in result:
                result[link.object_id].append(document)
    return result


def _evidence_table(documents, styles):
    if not documents:
        return None
    data = [[_paragraph(label, styles["cell_bold"]) for label in ("Title", "Type", "Uploaded", "Confidentiality")]]
    for document in sorted(documents, key=lambda item: (item.title.lower(), item.pk)):
        data.append(
            [
                _paragraph(document.title, styles["cell"]),
                _paragraph(document.document_type, styles["cell"]),
                _paragraph(timezone.localtime(document.uploaded_at).strftime("%d %b %Y"), styles["cell"]),
                _paragraph(document.get_confidentiality_level_display(), styles["cell"]),
            ]
        )
    return _table(data, (82 * mm, 35 * mm, 30 * mm, 28 * mm))


def build_pdf_bytes(user, dataset):
    regular, bold = _register_pdf_fonts()
    styles = _pdf_styles(regular, bold)
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=22 * mm,
        title=PDF_TITLE,
        author="TNK Insight",
    )
    story = [
        Paragraph(PDF_TITLE, styles["title"]),
        Paragraph(
            "Scoped report metadata only. Restricted personal, household, health, disability, safety, and financial details are excluded.",
            styles["subtitle"],
        ),
        Spacer(1, 5 * mm),
        _filter_summary(dataset, styles),
        Spacer(1, 6 * mm),
        Paragraph("Report summary", styles["heading"]),
        _summary_table(dataset, styles),
    ]
    evidence = _evidence_for_reports(user, dataset.reports)
    detailed_reports = [report for report in dataset.reports if can_view_report(user, report)]
    for report in detailed_reports:
        story.extend(
            [
                PageBreak(),
                Paragraph(f"{_safe_text(report.village.name_en)} - {_safe_text(report.reporting_period)}", styles["heading"]),
                _report_metadata(report, styles),
                Spacer(1, 5 * mm),
                Paragraph("Approval history", styles["subheading"]),
                _approval_table(report, styles),
            ]
        )
        projects = _project_table(user, report, styles)
        if projects is not None:
            story.extend([Spacer(1, 5 * mm), Paragraph("IVDP project summary", styles["subheading"]), projects])
        amendments = _amendment_table(report, styles)
        if amendments is not None:
            story.extend([Spacer(1, 5 * mm), Paragraph("Approved amendments", styles["subheading"]), amendments])
        evidence_table = _evidence_table(evidence.get(report.pk, ()), styles)
        if evidence_table is not None:
            story.extend([Spacer(1, 5 * mm), Paragraph("Evidence references", styles["subheading"]), evidence_table])
        story.append(
            KeepTogether(
                [
                    Spacer(1, 5 * mm),
                    Paragraph(
                        "This export is a presentation of the authorised report record. Approved amendment overlays are authoritative; original official data remains preserved in TNK Insight.",
                        styles["small"],
                    ),
                ]
            )
        )
    document.build(story, canvasmaker=_pdf_canvasmaker())
    return output.getvalue()


def pdf_export(user, reports, filters, reason):
    dataset = build_export_dataset(user, reports, filters, reason)
    content = build_pdf_bytes(user, dataset)
    audit(user, "pdf", len(dataset.rows), dataset.filters, dataset.reason)
    return _response(content, "application/pdf", "tnk-reports.pdf")
